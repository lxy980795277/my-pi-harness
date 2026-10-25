"""Native presentation objects, with explicit template and layout checks."""
from common import audit_package, read_package, require


def inspect_pptx(path):
    from pptx import Presentation
    parts = read_package(path)
    presentation = Presentation(path)
    require(0 < len(presentation.slides) <= 100, 'Presentation needs 1–100 slides')
    slides = []
    for slide_index, slide in enumerate(presentation.slides):
        shapes = []
        for shape_index, shape in enumerate(slide.shapes):
            require(shape.left >= 0 and shape.top >= 0 and shape.left + shape.width <= presentation.slide_width + 1 and shape.top + shape.height <= presentation.slide_height + 1, 'Shape exceeds slide bounds at slide ' + str(slide_index) + ', shape ' + str(shape_index))
            item = {'index': shape_index, 'name': shape.name}
            if shape.has_text_frame:
                item['text'] = shape.text[:20000]
            if shape.has_table:
                require(len(shape.table.rows) * len(shape.table.columns) <= 20000, 'Presentation table exceeds cell limit')
                item['table'] = [[cell.text for cell in row.cells] for row in shape.table.rows]
            if shape.has_chart:
                series = shape.chart.series
                require(len(series) > 0, 'Chart has no series')
                require(all(len(item.values) > 0 for item in series), 'Chart has an empty series')
                item['chart'] = [{'name': series_item.name, 'values': list(series_item.values)} for series_item in series]
            shapes.append(item)
        slides.append({'index': slide_index, 'shapes': shapes})
    return {'format': 'pptx', 'widthEmu': presentation.slide_width, 'heightEmu': presentation.slide_height, 'slides': slides,
            'validation': audit_package(parts) + ['Native shapes/tables/chart series inspected', 'No shape exceeds slide bounds'],
            'notVerified': ['PowerPoint/LibreOffice slide rendering', 'Text fit and visual overlap require rendered inspection']}


def category_data(spec):
    from pptx.chart.data import CategoryChartData
    require(all(len(item['values']) == len(spec['categories']) for item in spec['series']), 'Chart categories and values must have equal lengths')
    require(spec['kind'] != 'pie' or len(spec['series']) == 1, 'Pie chart requires one series')
    data = CategoryChartData()
    data.categories = spec['categories']
    for series in spec['series']:
        data.add_series(series['name'], series['values'])
    return data


def pptx(request, output):
    from pptx import Presentation
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt
    source = request.get('source')
    original_parts = read_package(source) if source else None
    presentation = Presentation(source) if source else Presentation()
    if source:
        require('width' not in request and 'height' not in request, 'Template slide dimensions cannot be changed')
    else:
        presentation.slide_width = Inches(request.get('width', 13.333333))
        presentation.slide_height = Inches(request.get('height', 7.5))
    original_size = (presentation.slide_width, presentation.slide_height)
    for slide_spec in request.get('slides', []):
        layout = slide_spec.get('layout', min(6, len(presentation.slide_layouts) - 1))
        require(layout < len(presentation.slide_layouts), 'Template slide layout does not exist')
        slide = presentation.slides.add_slide(presentation.slide_layouts[layout])
        for spec in slide_spec['shapes']:
            x, y, width, height = [Inches(spec[key]) for key in ('x', 'y', 'width', 'height')]
            require(x + width <= presentation.slide_width + 1 and y + height <= presentation.slide_height + 1, 'New shape exceeds slide bounds')
            if spec['kind'] == 'text':
                shape = slide.shapes.add_textbox(x, y, width, height)
                shape.text_frame.text = spec['text']
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(spec.get('fontSize', 20))
            elif spec['kind'] == 'table':
                rows = spec['rows']
                require(rows and rows[0] and all(len(row) == len(rows[0]) for row in rows), 'Slide table must be rectangular')
                shape = slide.shapes.add_table(len(rows), len(rows[0]), x, y, width, height)
                for row_index, row in enumerate(rows):
                    for column, value in enumerate(row):
                        shape.table.cell(row_index, column).text = '' if value is None else str(value)
            else:
                chart_type = {'bar': XL_CHART_TYPE.COLUMN_CLUSTERED, 'line': XL_CHART_TYPE.LINE, 'pie': XL_CHART_TYPE.PIE}[spec['chart']['kind']]
                shape = slide.shapes.add_chart(chart_type, x, y, width, height, category_data(spec['chart']))
                if 'title' in spec['chart']:
                    shape.chart.has_title = True
                    shape.chart.chart_title.text_frame.text = spec['chart']['title']
    for edit in request.get('edits', []):
        require(edit['slide'] < len(presentation.slides), 'Edited slide does not exist')
        slide = presentation.slides[edit['slide']]
        require(edit['shape'] < len(slide.shapes), 'Edited shape does not exist')
        shape = slide.shapes[edit['shape']]
        if edit['kind'] == 'text':
            require(shape.has_text_frame, 'Selected shape has no editable text frame')
            shape.text_frame.text = edit['text']
        elif edit['kind'] == 'table_cell':
            require(shape.has_table, 'Selected shape is not a native table')
            require(edit['row'] < len(shape.table.rows) and edit['column'] < len(shape.table.columns), 'Edited table cell does not exist')
            shape.table.cell(edit['row'], edit['column']).text = edit['text']
        else:
            require(shape.has_chart, 'Selected shape is not a native chart')
            expected_type = {'bar': XL_CHART_TYPE.COLUMN_CLUSTERED, 'line': XL_CHART_TYPE.LINE, 'pie': XL_CHART_TYPE.PIE}[edit['chart']['kind']]
            require(shape.chart.chart_type == expected_type, 'Changing an existing chart type requires creating a new chart')
            shape.chart.replace_data(category_data(edit['chart']))
    require(len(presentation.slides) <= 100, 'Presentation slide count exceeds limit')
    presentation.save(output)
    if original_parts:
        saved = read_package(output)
        require(set(original_parts).issubset(saved), 'Template package parts were removed during save')
        reopened = Presentation(output)
        require((reopened.slide_width, reopened.slide_height) == original_size, 'Template slide dimensions changed')
    result = inspect_pptx(output)
    if source:
        result['validation'].append('Template part names and slide dimensions preserved')
    return result
