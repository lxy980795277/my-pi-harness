"""Restricted DOCX package edits preserve unrelated OOXML parts."""
import re
from lxml import etree as ET
from common import NS, audit_package, clone, ensure_part, paragraph_text, plain_paragraph, read_package, require, set_xml, tag, write_package, xml

W = NS['w']
NOTE_CONTENT = 'application/vnd.openxmlformats-officedocument.wordprocessingml.'


def inspect_docx(path):
    parts = read_package(path)
    doc = xml(parts['word/document.xml'])
    paragraphs = list(doc.iter(tag('w', 'p')))
    starts = [node.get(tag('w', 'id')) for node in doc.iter(tag('w', 'bookmarkStart'))]
    ends = [node.get(tag('w', 'id')) for node in doc.iter(tag('w', 'bookmarkEnd'))]
    require(len(starts) == len(set(starts)) and sorted(starts) == sorted(ends), 'Bookmark start/end references are inconsistent')
    for kind in ('footnote', 'endnote', 'comment'):
        references = [node.get(tag('w', 'id')) for node in doc.iter(tag('w', kind + 'Reference'))]
        if not references:
            continue
        name = 'word/' + kind + 's.xml'
        require(name in parts, 'Referenced note/comment part is missing')
        identifiers = [node.get(tag('w', 'id')) for node in xml(parts[name])]
        require(len(identifiers) == len(set(identifiers)) and set(references).issubset(identifiers), 'Note/comment reference is missing or duplicated')
    return {
        'format': 'docx', 'paragraphs': len(paragraphs),
        'tables': len(list(doc.iter(tag('w', 'tbl')))),
        'text': '\n'.join(paragraph_text(p) for p in paragraphs)[:100000],
        'revisions': sum(len(list(doc.iter(tag('w', name)))) for name in ('ins', 'del')),
        'comments': len(list(xml(parts['word/comments.xml']))) if 'word/comments.xml' in parts else 0,
        'bookmarks': [node.get(tag('w', 'name')) for node in doc.iter(tag('w', 'bookmarkStart'))],
        'fields': [node.text for node in doc.iter(tag('w', 'instrText'))],
        'contentControls': len(list(doc.iter(tag('w', 'sdt')))),
        'footnotes': 'word/footnotes.xml' in parts, 'endnotes': 'word/endnotes.xml' in parts,
        'validation': audit_package(parts),
        'notVerified': ['Word/LibreOffice page rendering', 'Field/TOC calculation; fields require an Office application update'],
    }


def edit_revisions(doc, action):
    forbidden = ('moveFrom', 'moveTo', 'rPrChange', 'pPrChange', 'tblPrChange', 'trPrChange', 'tcPrChange', 'sectPrChange', 'cellIns', 'cellDel', 'cellMerge')
    require(not any(list(doc.iter(tag('w', name))) for name in forbidden), 'Complex revisions require a dedicated review workflow')
    parents = {child: parent for parent in doc.iter() for child in parent}
    for node in list(doc.iter()):
        if node.tag not in (tag('w', 'ins'), tag('w', 'del')):
            continue
        parent = parents[node]
        require(parent.tag == tag('w', 'p'), 'Only plain paragraph insertion/deletion revisions are supported')
        require(all(child.tag == tag('w', 'r') for child in node), 'Nested/structural revisions are unsupported')
        require(all(grandchild.tag in (tag('w', 'rPr'), tag('w', 't'), tag('w', 'delText'), tag('w', 'tab'), tag('w', 'br')) for child in node for grandchild in child), 'Revision contains unsupported run content')
        keep = (node.tag == tag('w', 'ins')) == (action == 'accept')
        offset = list(parent).index(node)
        if keep:
            for child in list(node):
                for deleted in child.iter(tag('w', 'delText')):
                    deleted.tag = tag('w', 't')
                parent.insert(offset, child)
                offset += 1
        parent.remove(node)


def replace_text(doc, old, replacement):
    matches = [p for p in doc.iter(tag('w', 'p')) if old in paragraph_text(p)]
    require(len(matches) == 1 and paragraph_text(matches[0]).count(old) == 1, 'Exact text must occur once in one paragraph')
    paragraph = matches[0]
    require(not any(list(paragraph.iter(tag('w', name))) for name in ('fldChar', 'instrText', 'hyperlink', 'ins', 'del', 'drawing', 'sdt')), 'Replacement crosses a field, hyperlink, revision or complex object')
    nodes = list(paragraph.iter(tag('w', 't')))
    start = paragraph_text(paragraph).index(old)
    end = start + len(old)
    cursor = 0
    inserted = False
    for node in nodes:
        original = node.text or ''
        left, right = cursor, cursor + len(original)
        cursor = right
        if right <= start or left >= end:
            continue
        before = original[:max(0, start - left)]
        after = original[max(0, end - left):]
        node.text = before + (replacement if not inserted else '') + after
        node.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        inserted = True


def add_note(parts, paragraph, operation):
    kind = operation['noteType']
    plural = kind + 's'
    name = 'word/' + plural + '.xml'
    root = ensure_part(parts, name, plural, NOTE_CONTENT + plural + '+xml', plural)
    if not list(root):
        for identifier, note_type, separator in ((-1, 'separator', 'separator'), (0, 'continuationSeparator', 'continuationSeparator')):
            note = ET.SubElement(root, tag('w', kind), {tag('w', 'id'): str(identifier), tag('w', 'type'): note_type})
            run = ET.SubElement(ET.SubElement(note, tag('w', 'p')), tag('w', 'r'))
            ET.SubElement(run, tag('w', separator))
    identifier = max([int(node.get(tag('w', 'id'), '0')) for node in root] + [0]) + 1
    note = ET.SubElement(root, tag('w', kind), {tag('w', 'id'): str(identifier)})
    note.append(plain_paragraph(operation['text']))
    run = ET.SubElement(paragraph, tag('w', 'r'))
    ET.SubElement(run, tag('w', kind + 'Reference'), {tag('w', 'id'): str(identifier)})
    set_xml(parts, name, root)


def insert_paragraph_start(paragraph, node):
    offset = 1 if len(paragraph) and paragraph[0].tag == tag('w', 'pPr') else 0
    paragraph.insert(offset, node)


def comment(parts, doc, paragraphs, operation):
    root = ensure_part(parts, 'word/comments.xml', 'comments', NOTE_CONTENT + 'comments+xml', 'comments')
    if operation['kind'] == 'comment_add':
        paragraph = paragraphs[operation['paragraph']]
        require(len(list(paragraph.iter(tag('w', 'r')))) > 0, 'A comment needs a text run')
        identifier = max([int(node.get(tag('w', 'id'), '0')) for node in root] + [-1]) + 1
        node = ET.SubElement(root, tag('w', 'comment'), {tag('w', 'id'): str(identifier), tag('w', 'author'): operation['author'], tag('w', 'initials'): operation['author'][:3]})
        node.append(plain_paragraph(operation['text']))
        insert_paragraph_start(paragraph, ET.Element(tag('w', 'commentRangeStart'), {tag('w', 'id'): str(identifier)}))
        ET.SubElement(paragraph, tag('w', 'commentRangeEnd'), {tag('w', 'id'): str(identifier)})
        ET.SubElement(ET.SubElement(paragraph, tag('w', 'r')), tag('w', 'commentReference'), {tag('w', 'id'): str(identifier)})
    else:
        identifier = str(operation['id'])
        found = [node for node in root if node.get(tag('w', 'id')) == identifier]
        require(len(found) == 1, 'Comment id does not exist')
        root.remove(found[0])
        for parent in doc.iter():
            for child in list(parent):
                if child.tag in {tag('w', 'commentRangeStart'), tag('w', 'commentRangeEnd'), tag('w', 'commentReference')} and child.get(tag('w', 'id')) == identifier:
                    parent.remove(child)
    set_xml(parts, 'word/comments.xml', root)


def append_body(body, node):
    sections = body.find(tag('w', 'sectPr'))
    if sections is None:
        body.append(node)
    else:
        body.insert(list(body).index(sections), node)


def add_field(body, doc, operation):
    kind = operation['fieldType']
    target = operation.get('target', '')
    if kind != 'toc':
        require(re.fullmatch(r'[A-Za-z_][A-Za-z0-9_]{0,39}', target) is not None, 'Field target must be a safe bookmark or sequence name')
    if kind == 'ref':
        require(any(node.get(tag('w', 'name')) == target for node in doc.iter(tag('w', 'bookmarkStart'))), 'REF bookmark does not exist')
    instruction = {'toc': ' TOC \\o "1-3" \\h \\z \\u ', 'ref': ' REF ' + target + ' \\h ', 'caption': ' SEQ ' + target + ' \\* ARABIC '}[kind]
    paragraph = ET.Element(tag('w', 'p'))
    if kind == 'caption':
        ET.SubElement(ET.SubElement(paragraph, tag('w', 'r')), tag('w', 't')).text = target + ' '
    for char_type in ('begin', 'separate', 'end'):
        run = ET.SubElement(paragraph, tag('w', 'r'))
        ET.SubElement(run, tag('w', 'fldChar'), {tag('w', 'fldCharType'): char_type, tag('w', 'dirty'): 'true'})
        if char_type == 'begin':
            ET.SubElement(ET.SubElement(paragraph, tag('w', 'r')), tag('w', 'instrText')).text = instruction
        if char_type == 'separate':
            ET.SubElement(ET.SubElement(paragraph, tag('w', 'r')), tag('w', 't')).text = ('[Update field in Word]' if kind == 'caption' else operation.get('text', '[Update field in Word]'))
    if kind == 'caption' and operation.get('text'):
        ET.SubElement(ET.SubElement(paragraph, tag('w', 'r')), tag('w', 't')).text = ' ' + operation['text']
    append_body(body, paragraph)


def merge(parts, body, paths):
    if not paths:
        return
    require(not any(name in parts for name in ('word/comments.xml', 'word/footnotes.xml', 'word/endnotes.xml')), 'Merge requires simple documents without comments/notes')
    forbidden = ('drawing', 'pict', 'numPr', 'bookmarkStart', 'footnoteReference', 'endnoteReference', 'commentReference', 'hyperlink', 'ins', 'del', 'fldChar', 'sdt', 'altChunk')
    for path in paths:
        other_parts = read_package(path)
        require(not any(name in other_parts for name in ('word/comments.xml', 'word/footnotes.xml', 'word/endnotes.xml')), 'Merge source has notes/comments')
        other = xml(other_parts['word/document.xml'])
        require(not any(list(other.iter(tag('w', name))) for name in forbidden), 'Merge source contains complex or related structures')
        styles = {node.get(tag('w', 'styleId')) for node in xml(parts['word/styles.xml']).iter(tag('w', 'style'))}
        for name in ('pStyle', 'rStyle', 'tblStyle'):
            require(all(node.get(tag('w', 'val')) in styles for node in other.iter(tag('w', name))), 'Merge source references a style absent from the base document')
        other_body = other.find(tag('w', 'body'))
        for node in other_body:
            if node.tag == tag('w', 'sectPr'):
                continue
            require(node.tag in (tag('w', 'p'), tag('w', 'tbl')), 'Unsupported merge body structure')
            append_body(body, clone(node))


def docx(request, output):
    from docx import Document
    source = request.get('source')
    if source:
        require(not request.get('paragraphs') and not request.get('tables'), 'Source edits use explicit operations; creation paragraphs/tables cannot silently replace content')
        parts = read_package(source)
    else:
        document = Document()
        for item in request.get('paragraphs', []):
            if 'heading' in item:
                document.add_heading(item['text'], item['heading'])
            else:
                document.add_paragraph(item['text'])
        for table in request.get('tables', []):
            rows = table['rows']
            require(rows and rows[0] and all(len(row) == len(rows[0]) for row in rows), 'A table must be rectangular')
            result = document.add_table(rows=len(rows), cols=len(rows[0]))
            for row_index, row in enumerate(rows):
                for column, value in enumerate(row):
                    result.cell(row_index, column).text = '' if value is None else str(value)
        document.save(output)
        parts = read_package(output)
    doc = xml(parts['word/document.xml'])
    body = doc.find(tag('w', 'body'))
    require(body is not None, 'DOCX body is missing')
    for operation in request.get('operations', []):
        kind = operation['kind']
        paragraphs = list(body.iter(tag('w', 'p')))
        if 'paragraph' in operation:
            require(operation['paragraph'] < len(paragraphs), 'Paragraph index is outside the document')
        if kind == 'replace_text':
            replace_text(doc, operation['old'], operation['replacement'])
        elif kind == 'revisions':
            edit_revisions(doc, operation['action'])
        elif kind in ('comment_add', 'comment_remove'):
            comment(parts, doc, paragraphs, operation)
        elif kind == 'note':
            add_note(parts, paragraphs[operation['paragraph']], operation)
        elif kind == 'bookmark':
            require(not any(node.get(tag('w', 'name')) == operation['name'] for node in doc.iter(tag('w', 'bookmarkStart'))), 'Bookmark already exists')
            identifier = str(max([int(node.get(tag('w', 'id'), '0')) for node in doc.iter(tag('w', 'bookmarkStart'))] + [0]) + 1)
            paragraph = paragraphs[operation['paragraph']]
            insert_paragraph_start(paragraph, ET.Element(tag('w', 'bookmarkStart'), {tag('w', 'id'): identifier, tag('w', 'name'): operation['name']}))
            ET.SubElement(paragraph, tag('w', 'bookmarkEnd'), {tag('w', 'id'): identifier})
        elif kind == 'field':
            add_field(body, doc, operation)
        elif kind == 'content_control':
            node = ET.Element(tag('w', 'sdt'))
            props = ET.SubElement(node, tag('w', 'sdtPr'))
            ET.SubElement(props, tag('w', 'tag'), {tag('w', 'val'): operation['tag']})
            ET.SubElement(props, tag('w', 'text'))
            ET.SubElement(node, tag('w', 'sdtContent')).append(plain_paragraph(operation['text']))
            append_body(body, node)
        elif kind == 'metadata':
            core = xml(parts['docProps/core.xml'])
            uris = {'title': 'http://purl.org/dc/elements/1.1/', 'subject': 'http://purl.org/dc/elements/1.1/', 'creator': 'http://purl.org/dc/elements/1.1/', 'description': 'http://purl.org/dc/elements/1.1/'}
            for name, uri in uris.items():
                if name not in operation:
                    continue
                key = '{' + uri + '}' + name
                node = core.find(key)
                if node is None:
                    node = ET.SubElement(core, key)
                node.text = operation[name]
            set_xml(parts, 'docProps/core.xml', core)
    merge(parts, body, request.get('mergeSources', []))
    set_xml(parts, 'word/document.xml', doc)
    write_package(parts, output)
    result = inspect_docx(output)
    result['validation'].append('Output reopened; original input not overwritten')
    return result
