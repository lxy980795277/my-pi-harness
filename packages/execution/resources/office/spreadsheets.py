"""Native spreadsheet editing; formula caches are inspected, never invented."""
import re
from common import audit_package, read_package, require, xml

S = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


def inspect_xlsx(path):
    from openpyxl import load_workbook
    parts = read_package(path)
    require(not any(name.startswith('xl/externalLinks/') for name in parts), 'External workbook links are not supported')
    workbook = load_workbook(path, data_only=False, keep_links=False)
    require(len(workbook.worksheets) <= 20, 'Workbook has too many sheets')
    sheets, errors, formulas, caches = [], [], [], []
    total = 0
    for sheet in workbook.worksheets:
        require(sheet.max_row <= 20000 and sheet.max_column <= 1000, 'Worksheet dimensions exceed limit')
        total += sheet.max_row * sheet.max_column
        require(total <= 20000, 'Workbook cell area exceeds limit')
        rows = []
        for row in sheet.iter_rows():
            values = []
            for cell in row:
                if cell.data_type == 'e':
                    errors.append({'sheet': sheet.title, 'cell': cell.coordinate, 'value': cell.value})
                if cell.data_type == 'f':
                    formulas.append({'sheet': sheet.title, 'cell': cell.coordinate, 'formula': cell.value})
                values.append(cell.value)
            rows.append(values)
        sheets.append({'name': sheet.title, 'rows': rows[:200], 'rowCount': sheet.max_row, 'columnCount': sheet.max_column, 'tables': list(sheet.tables)})
    for name, data in parts.items():
        if name.startswith('xl/worksheets/sheet') and name.endswith('.xml'):
            for cell in xml(data).iter(S + 'c'):
                if cell.find(S + 'f') is not None:
                    value = cell.find(S + 'v')
                    caches.append({'part': name, 'cell': cell.get('r'), 'cachedValue': value.text if value is not None else None})
    workbook.close()
    missing = sum(item['cachedValue'] is None for item in caches)
    not_verified = ['Excel/LibreOffice rendering', 'Formulas have not been recalculated by this tool']
    if missing:
        not_verified.append(str(missing) + ' formula cache(s) are absent; expected numerical results cannot be verified')
    return {'format': 'xlsx', 'sheets': sheets, 'formulas': formulas[:1000], 'formulaCaches': caches[:1000], 'errorCells': errors[:1000],
            'charts': sum(name.startswith('xl/charts/chart') and name.endswith('.xml') for name in parts),
            'validation': audit_package(parts) + ['Workbook reopened; formulas and cached values inspected independently'],
            'notVerified': not_verified}


def safe_value(value):
    if isinstance(value, str) and value.startswith('='):
        require(not re.search(r'\[|\]|https?://|file:|WEBSERVICE\s*\(|HYPERLINK\s*\(|RTD\s*\(|\b[A-Za-z]+\|', value, re.I), 'External/network/DDE formulas are unsupported')
        require('#REF!' not in value, 'Formula contains a broken reference')
    return value


def chart_data(workbook, request, number):
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    title = '_office_chart_' + str(number)
    require(title not in workbook.sheetnames, 'Chart data sheet name conflict')
    require(len(workbook.worksheets) < 20, 'Chart data exceeds worksheet limit')
    sheet = workbook.create_sheet(title)
    sheet.sheet_state = 'hidden'
    spec = request['chart']
    categories = spec['categories']
    require(all(len(series['values']) == len(categories) for series in spec['series']), 'Chart categories and values must have equal length')
    require(spec['kind'] != 'pie' or len(spec['series']) == 1, 'Pie charts require one series')
    sheet.append(['Category'] + [series['name'] for series in spec['series']])
    for index, category in enumerate(categories):
        sheet.append([category] + [series['values'][index] for series in spec['series']])
    result = {'bar': BarChart, 'line': LineChart, 'pie': PieChart}[spec['kind']]()
    result.title = spec.get('title', 'Chart')
    result.add_data(Reference(sheet, min_col=2, max_col=len(spec['series']) + 1, min_row=1, max_row=len(categories) + 1), titles_from_data=True)
    result.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=len(categories) + 1))
    return result


def xlsx(request, output):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries
    source = request.get('source')
    if source:
        parts = read_package(source)
        require(not any(name.startswith('xl/externalLinks/') or name.startswith('xl/drawings/vml') or name.startswith('xl/slicer') for name in parts), 'External links, VML and slicers require a dedicated workflow')
        workbook = load_workbook(source, keep_links=False)
    else:
        workbook = Workbook()
        workbook.remove(workbook.active)
    chart_number = 0
    try:
        for spec in request['sheets']:
            require(not any(char in spec['name'] for char in '[]:*?/\\'), 'Invalid worksheet name')
            sheet = workbook[spec['name']] if spec['name'] in workbook.sheetnames else workbook.create_sheet(spec['name'])
            row_start, col_start = coordinate_to_tuple(spec.get('startCell', 'A1'))
            require(row_start <= 20000 and col_start <= 1000, 'Start cell is outside supported bounds')
            for row_offset, values in enumerate(spec.get('rows', [])):
                require(row_start + row_offset <= 20000 and col_start + len(values) <= 1001, 'Data exceeds supported worksheet bounds')
                for column_offset, value in enumerate(values):
                    sheet.cell(row_start + row_offset, col_start + column_offset, safe_value(value))
            for edit in spec.get('cells', []):
                row, column = coordinate_to_tuple(edit['cell'])
                require(row <= 20000 and column <= 1000, 'Edited cell is outside supported bounds')
                cell = sheet[edit['cell']]
                cell.value = safe_value(edit['value'])
                if 'bold' in edit or 'fontSize' in edit:
                    from copy import copy
                    font = copy(cell.font)
                    if 'bold' in edit:
                        font.bold = edit['bold']
                    if 'fontSize' in edit:
                        font.sz = edit['fontSize']
                    cell.font = font
                if 'fill' in edit:
                    cell.fill = PatternFill('solid', fgColor=edit['fill'])
                if 'numberFormat' in edit:
                    cell.number_format = edit['numberFormat']
            for table in spec.get('tables', []):
                minimum_col, minimum_row, maximum_col, maximum_row = range_boundaries(table['range'])
                require(maximum_row <= sheet.max_row and maximum_col <= sheet.max_column and maximum_row > minimum_row, 'Table range must contain an existing header and data')
                require(all(isinstance(sheet.cell(minimum_row, column).value, str) and sheet.cell(minimum_row, column).value for column in range(minimum_col, maximum_col + 1)), 'Table headers must be non-empty strings')
                if table['name'] in sheet.tables:
                    del sheet.tables[table['name']]
                result = Table(displayName=table['name'], ref=table['range'])
                result.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
                sheet.add_table(result)
            for chart in spec.get('charts', []):
                while '_office_chart_' + str(chart_number) in workbook.sheetnames:
                    chart_number += 1
                sheet.add_chart(chart_data(workbook, chart, chart_number), chart['anchor'])
                chart_number += 1
        require(len(workbook.worksheets) <= 20, 'Workbook has too many sheets')
        require(sum(sheet.max_row * sheet.max_column for sheet in workbook.worksheets) <= 20000, 'Workbook cell area exceeds limit')
        workbook.save(output)
    finally:
        workbook.close()
    result = inspect_xlsx(output)
    require(not result['errorCells'], 'Output contains spreadsheet error cells')
    return result
