"""Explicit DOCX typography and section settings without rebuilding existing packages."""
from lxml import etree as ET
from common import ensure_part, plain_paragraph, require, set_xml, tag, xml


def child(parent, name):
    node = parent.find(tag('w', name))
    if node is None:
        node = ET.SubElement(parent, tag('w', name))
    return node


def properties(parent, name):
    node = parent.find(tag('w', name))
    if node is None:
        node = ET.Element(tag('w', name))
        parent.insert(0, node)
    return node


def font_properties(props, values):
    if 'name' in values:
        node = child(props, 'rFonts')
        for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            node.set(tag('w', key), values['name'])
        for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            node.attrib.pop(tag('w', key), None)
    if 'size' in values:
        for key in ('sz', 'szCs'):
            child(props, key).set(tag('w', 'val'), str(round(values['size'] * 2)))
    if 'color' in values:
        child(props, 'color').set(tag('w', 'val'), values['color'].upper())
    for option, key in (('bold', 'b'), ('italic', 'i'), ('underline', 'u')):
        if option in values:
            value = ('single' if values[option] else 'none') if key == 'u' else ('1' if values[option] else '0')
            child(props, key).set(tag('w', 'val'), value)


def paragraph_properties(props, values):
    if 'alignment' in values:
        child(props, 'jc').set(tag('w', 'val'), 'both' if values['alignment'] == 'justify' else values['alignment'])
    for option, attribute in (('spaceBefore', 'before'), ('spaceAfter', 'after')):
        if option in values:
            child(props, 'spacing').set(tag('w', attribute), str(round(values[option] * 20)))
    if 'lineSpacing' in values:
        spacing = child(props, 'spacing')
        spacing.set(tag('w', 'line'), str(round(values['lineSpacing'] * 240)))
        spacing.set(tag('w', 'lineRule'), 'auto')
    for option, attribute in (('leftIndent', 'left'), ('rightIndent', 'right'), ('firstLineIndent', 'firstLine')):
        if option in values:
            indent = child(props, 'ind')
            if option == 'firstLineIndent':
                indent.attrib.pop(tag('w', 'hanging'), None)
            amount = values[option]
            if option == 'firstLineIndent' and amount < 0:
                attribute, amount = 'hanging', -amount
                indent.attrib.pop(tag('w', 'firstLine'), None)
            indent.set(tag('w', attribute), str(round(amount * 20)))
    for option, key in (('keepWithNext', 'keepNext'), ('keepTogether', 'keepLines'), ('pageBreakBefore', 'pageBreakBefore')):
        if option in values:
            child(props, key).set(tag('w', 'val'), '1' if values[option] else '0')


def style_identifier(styles, name):
    for node in styles.findall(tag('w', 'style')):
        display = node.find(tag('w', 'name'))
        if node.get(tag('w', 'styleId')) == name or (display is not None and display.get(tag('w', 'val')) == name):
            require(node.get(tag('w', 'type')) == 'paragraph', 'A paragraph requires a paragraph style')
            return node.get(tag('w', 'styleId'))
    raise ValueError('Paragraph style does not exist: ' + name)


def format_paragraph(paragraph, values, styles):
    if 'style' in values:
        child(properties(paragraph, 'pPr'), 'pStyle').set(tag('w', 'val'), style_identifier(styles, values['style']))
    if 'format' in values:
        paragraph_properties(properties(paragraph, 'pPr'), values['format'])
    if 'font' in values:
        for run in paragraph.iter(tag('w', 'r')):
            font_properties(properties(run, 'rPr'), values['font'])


def apply_styles(parts, request):
    styles = xml(parts['word/styles.xml'])
    if 'defaultFont' in request:
        normal = next(node for node in styles.findall(tag('w', 'style')) if node.get(tag('w', 'styleId')) == 'Normal')
        font_properties(child(normal, 'rPr'), request['defaultFont'])
    names = [entry['name'] for entry in request.get('styles', [])]
    require(len(names) == len(set(names)), 'Style names must be unique')
    for index, entry in enumerate(request.get('styles', [])):
        existing = [node for node in styles.findall(tag('w', 'style')) if node.find(tag('w', 'name')) is not None and node.find(tag('w', 'name')).get(tag('w', 'val')) == entry['name']]
        if existing:
            style = existing[0]
            require(style.get(tag('w', 'type')) == 'paragraph', 'Only paragraph styles can be edited')
        else:
            identifier = 'MyPiStyle' + str(index)
            known = {node.get(tag('w', 'styleId')) for node in styles}
            while identifier in known:
                identifier += 'x'
            style = ET.SubElement(styles, tag('w', 'style'), {tag('w', 'type'): 'paragraph', tag('w', 'styleId'): identifier})
            child(style, 'name').set(tag('w', 'val'), entry['name'])
        if not existing or 'basedOn' in entry:
            base = style_identifier(styles, entry.get('basedOn', 'Normal'))
            require(base != style.get(tag('w', 'styleId')), 'A style cannot inherit itself')
            child(style, 'basedOn').set(tag('w', 'val'), base)
        if 'font' in entry:
            font_properties(child(style, 'rPr'), entry['font'])
        if 'format' in entry:
            paragraph_properties(child(style, 'pPr'), entry['format'])
    parents = {node.get(tag('w', 'styleId')): node.find(tag('w', 'basedOn')).get(tag('w', 'val'))
               for node in styles if node.find(tag('w', 'basedOn')) is not None}
    for name in names:
        current, visited = style_identifier(styles, name), set()
        while current is not None:
            require(current not in visited, 'Style inheritance cannot form a cycle')
            visited.add(current)
            current = parents.get(current)
    set_xml(parts, 'word/styles.xml', styles)
    return styles


def header_footer(parts, sections, kind, values):
    root_name = 'hdr' if kind == 'header' else 'ftr'
    number = 1
    while f'word/{kind}{number}.xml' in parts:
        number += 1
    name = f'word/{kind}{number}.xml'
    root = ensure_part(parts, name, root_name, f'application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}+xml', kind)
    paragraph = plain_paragraph(values['text'])
    if values.get('pageNumber'):
        field = ET.SubElement(paragraph, tag('w', 'fldSimple'), {tag('w', 'instr'): ' PAGE ', tag('w', 'dirty'): 'true'})
        ET.SubElement(ET.SubElement(field, tag('w', 'r')), tag('w', 't')).text = '1'
    paragraph_properties(properties(paragraph, 'pPr'), {'alignment': values.get('alignment', 'left')})
    for run in paragraph.iter(tag('w', 'r')):
        font_properties(properties(run, 'rPr'), values.get('font', {}))
    root.append(paragraph)
    set_xml(parts, name, root)
    relationships = xml(parts['word/_rels/document.xml.rels'])
    relation = next(node for node in relationships if node.get('Target') == name.split('/')[-1])
    for section in sections:
        for reference in list(section.findall(tag('w', kind + 'Reference'))):
            if reference.get(tag('w', 'type')) == 'default':
                section.remove(reference)
        section.insert(0, ET.Element(tag('w', kind + 'Reference'), {tag('w', 'type'): 'default', tag('r', 'id'): relation.get('Id')}))


def apply_settings(parts, doc, request):
    styles = apply_styles(parts, request)
    sections = list(doc.iter(tag('w', 'sectPr')))
    require(sections, 'DOCX section settings are missing')
    settings = request.get('layout', {})
    for section in sections:
        size = child(section, 'pgSz')
        for option, key in (('width', 'w'), ('height', 'h')):
            if option in settings:
                size.set(tag('w', key), str(round(settings[option] * 20)))
        width, height = int(size.get(tag('w', 'w'), '12240')), int(size.get(tag('w', 'h'), '15840'))
        size.set(tag('w', 'orient'), 'landscape' if width > height else 'portrait')
        margins = child(section, 'pgMar')
        for key, value in settings.get('margins', {}).items():
            margins.set(tag('w', key), str(round(value * 20)))
        for option, key in (('headerDistance', 'header'), ('footerDistance', 'footer')):
            if option in settings:
                margins.set(tag('w', key), str(round(settings[option] * 20)))
        require(width - int(margins.get(tag('w', 'left'), '1440')) - int(margins.get(tag('w', 'right'), '1440')) >= 1440, 'Page margins leave less than 72 points of usable width')
        require(height - int(margins.get(tag('w', 'top'), '1440')) - int(margins.get(tag('w', 'bottom'), '1440')) >= 1440, 'Page margins leave less than 72 points of usable height')
    for kind in ('header', 'footer'):
        if kind in request:
            header_footer(parts, sections, kind, request[kind])
    return styles


def create_document(request, output):
    from docx import Document
    from docx.shared import Pt
    document = Document()
    layout = request.get('layout', {})
    for section in document.sections:
        for key, attribute in (('width', 'page_width'), ('height', 'page_height'), ('headerDistance', 'header_distance'), ('footerDistance', 'footer_distance')):
            if key in layout:
                setattr(section, attribute, Pt(layout[key]))
        for key, value in layout.get('margins', {}).items():
            setattr(section, key + '_margin', Pt(value))
    for item in request.get('paragraphs', []):
        require(('text' in item) != ('runs' in item), 'A paragraph requires exactly one of text or runs')
        paragraph = document.add_heading('', item['heading']) if 'heading' in item else document.add_paragraph()
        if 'text' in item:
            paragraph.add_run(item['text'])
        else:
            for value in item['runs']:
                run = paragraph.add_run(value['text'])
                font_properties(properties(run._r, 'rPr'), value.get('font', {}))
    for item in request.get('tables', []):
        rows = item['rows']
        require(rows and rows[0] and all(len(row) == len(rows[0]) for row in rows), 'A table must be rectangular')
        table = document.add_table(rows=len(rows), cols=len(rows[0]))
        if 'style' in item:
            table.style = item['style']
        count = item.get('headerRows', 0)
        require(count <= len(rows), 'Repeating header count exceeds table rows')
        for row_index, row in enumerate(rows):
            if row_index < count:
                child(table.rows[row_index]._tr.get_or_add_trPr(), 'tblHeader')
            for column, value in enumerate(row):
                table.cell(row_index, column).text = '' if value is None else str(value)
    for item in request.get('images', []):
        from PIL import Image
        Image.MAX_IMAGE_PIXELS = 16_000_000
        with Image.open(item['source']) as image:
            require(image.format in ('PNG', 'JPEG') and image.width * image.height <= 16_000_000, 'Only bounded PNG/JPEG images can be embedded')
            height = item.get('height', item['width'] * image.height / image.width)
            image.verify()
        section = document.sections[-1]
        usable_width = section.page_width.pt - section.left_margin.pt - section.right_margin.pt
        usable_height = section.page_height.pt - section.top_margin.pt - section.bottom_margin.pt
        require(item['width'] <= usable_width and height <= usable_height, 'DOCX image exceeds the usable page area')
        picture = document.add_picture(item['source'], width=Pt(item['width']), height=Pt(item['height']) if 'height' in item else None)
        picture._inline.docPr.set('descr', item['altText'])
        if 'alignment' in item:
            paragraph_properties(properties(document.paragraphs[-1]._p, 'pPr'), {'alignment': item['alignment']})
    document.save(output)


def inspect_layout(parts, doc):
    sections = []
    for section in list(doc.iter(tag('w', 'sectPr')))[:50]:
        size, margins = section.find(tag('w', 'pgSz')), section.find(tag('w', 'pgMar'))
        sections.append({'size': {key: int(value) / 20 for key, value in ((ET.QName(k).localname, v) for k, v in size.attrib.items()) if key in ('w', 'h')} if size is not None else {},
                         'margins': {ET.QName(key).localname: int(value) / 20 for key, value in margins.attrib.items() if ET.QName(key).localname in ('top', 'right', 'bottom', 'left', 'header', 'footer')} if margins is not None else {}})
    return {'sections': sections,
            'imageCount': sum(name.startswith('word/media/') for name in parts),
            'headers': [name for name in parts if name.startswith('word/header') and name.endswith('.xml')],
            'footers': [name for name in parts if name.startswith('word/footer') and name.endswith('.xml')]}
